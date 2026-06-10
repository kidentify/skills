#!/usr/bin/env python3
"""Build a PP Tunas (Indonesia PP No. 17/2025) children's risk register as .xlsx.

This script exists so every run of the pp-tunas-risk-assessment skill produces
the same defensible, multi-sheet workbook structure — and so we don't reinvent
it each time. You supply the SCORED content as JSON; the script handles layout,
styling, severity colouring, and the determination banner.

USAGE
-----
    python build_risk_register.py --in data.json --out "PP_Tunas_Risk_Register.xlsx"

The values in `data.json` must come from your scored assessment, grounded in
live neimo. lookups. The script does NOT know any thresholds — it only formats
what you give it.

INPUT JSON SHAPE
----------------
{
  "product": {
    "name": "Acme Chat",
    "company": "Acme Inc.",
    "assessment_date": "2026-06-10",
    "assessor": "Jane Dev",
    "summary": "Free social app with public voice chat, open to all ages..."
  },
  "determination": {
    "result": "HIGH-RISK",                 # or "LOW-RISK"
    "minimum_serviceable_age": 16,          # int
    "driving_factors": ["a", "b", "f"],     # which factors pushed it high
    "rationale": "Public voice chat exposes children to unknown adults...",
    "next_step": "Gate under-16s out of voice chat or treat product as 16+."
  },
  "risk_factors": [
    {
      "id": "a",
      "factor": "Possibility of children interacting with unknown individuals",
      "present": true,
      "severity": "High",                   # None | Low | Medium | High
      "likelihood": "High",                 # None | Low | Medium | High
      "existing_mitigation": "None today.",
      "residual_risk": "High",
      "required_mitigation": "Age assurance before any stranger contact; ...",
      "citation": "PP No. 17/2025 Art. 5; Permen 9/2026 ..."
    }
    # ... one per statutory factor, typically a–g
  ],
  "feature_matrix": [
    {
      "feature": "Voice chat (public)",
      "min_age": "16",
      "threshold_age": "18",
      "assurance_required": "Yes",
      "kid_mechanism": "Session permission + AgeKit+",
      "status": "Ungated",
      "gap": "No age check before voice chat."
    }
    # ... one per feature (chat modalities, forums, username/avatar,
    #     profiling, monetization mechanics, ...)
  ],
  "obligations": [
    {
      "requirement": "Age assurance",
      "applies": "Yes",
      "current_status": "Not implemented",
      "gap": "No age check at signup.",
      "action": "Deploy age assurance commensurate with risk.",
      "kid_mechanism": "Age Gate + AgeKit+",
      "owner": "Eng",
      "citation": "PP No. 17/2025 Art. 22"
    }
    # ... one per discrete obligation — be exhaustive (20-35+ rows)
  ],
  "sources": [
    {"topic": "Risk factors / DPIA",
     "regulation": "PP No. 17/2025",
     "article": "Art. 5",
     "url": "https://peraturan.bpk.go.id/Details/316698/pp-no-17-tahun-2025",
     "verification": "kid-reviewed"}
    # ... every neimo. citation used
  ]
}

Requires openpyxl (pip install openpyxl --break-system-packages).
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---- palette -------------------------------------------------------------
NAVY = "1F2A44"
SLATE = "33415C"
PAPER = "F4F6FA"
WHITE = "FFFFFF"
HIGH = "F4CCCC"     # red-ish
MED = "FCE5CD"      # orange-ish
LOW = "FFF2CC"      # yellow-ish
NONE_ = "D9EAD3"    # green-ish
HIGH_BANNER = "CC0000"
LOW_BANNER = "38761D"

SEV_FILL = {"high": HIGH, "medium": MED, "med": MED, "low": LOW,
            "none": NONE_, "n/a": NONE_, "": WHITE}

THIN = Side(style="thin", color="C8CFDC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _hdr(cell):
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=SLATE)
    cell.alignment = CENTER
    cell.border = BORDER


def _sev_fill(val):
    return PatternFill("solid", fgColor=SEV_FILL.get(str(val).strip().lower(), WHITE))


def _autorow(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.border.left.style is None:
                c.border = BORDER
            if c.alignment.wrap_text is None:
                c.alignment = WRAP


def build(data, out_path):
    wb = Workbook()

    # ---- Summary --------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    p = data.get("product", {})
    d = data.get("determination", {})

    ws.merge_cells("A1:D1")
    t = ws["A1"]
    t.value = "PP Tunas Children's Risk Assessment — Indonesia (PP No. 17/2025 + Permen 9/2026)"
    t.font = Font(bold=True, color=WHITE, size=14)
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = CENTER
    ws.row_dimensions[1].height = 34

    meta = [
        ("Product", p.get("name", "")),
        ("Company / ESO", p.get("company", "")),
        ("Assessment date", p.get("assessment_date", "")),
        ("Assessor", p.get("assessor", "")),
        ("Product summary", p.get("summary", "")),
    ]
    r = 3
    for label, val in meta:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=PAPER)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 2, val).alignment = WRAP
        r += 1

    r += 1
    # Determination banner
    result = str(d.get("result", "")).upper()
    banner = HIGH_BANNER if "HIGH" in result else LOW_BANNER if "LOW" in result else SLATE
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    b = ws.cell(r, 1, f"DETERMINATION: {result or 'NOT SET'}    "
                      f"Minimum serviceable age: {d.get('minimum_serviceable_age', 'TBD')}")
    b.font = Font(bold=True, color=WHITE, size=13)
    b.fill = PatternFill("solid", fgColor=banner)
    b.alignment = CENTER
    ws.row_dimensions[r].height = 28
    r += 2

    detail = [
        ("Driving factor(s)", ", ".join(d.get("driving_factors", [])) or "—"),
        ("Rationale", d.get("rationale", "")),
        ("What to do next", d.get("next_step", "")),
    ]
    for label, val in detail:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=PAPER)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 2, val).alignment = WRAP
        r += 1

    r += 1
    proviso = data.get("proviso") or (
        "AI-GENERATED — REVIEW REQUIRED. This risk assessment was generated by AI "
        "using neimo. regulatory data. It is a structured engineering aid, NOT legal "
        "advice, and must be reviewed by a qualified risk assessor and/or a qualified "
        "attorney before it is relied upon or filed. The regulator (Komdigi) makes the "
        "final HIGH/LOW-risk determination.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=4)
    note = ws.cell(r, 1, proviso)
    note.font = Font(bold=True, italic=True, color="9C0006", size=10)
    note.fill = PatternFill("solid", fgColor="FFF2CC")
    note.alignment = WRAP
    note.border = BORDER
    ws.row_dimensions[r].height = 30
    ws.column_dimensions["A"].width = 22
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 30

    # ---- Risk Factors ---------------------------------------------------
    ws = wb.create_sheet("Risk Factors")
    cols = ["#", "Mandatory risk factor", "Present?", "Severity", "Likelihood",
            "Existing mitigation", "Residual risk", "Required mitigation", "Citation"]
    for i, c in enumerate(cols, 1):
        _hdr(ws.cell(1, i, c))
    widths = [4, 34, 9, 11, 11, 30, 11, 36, 34]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    rr = 2
    for f in data.get("risk_factors", []):
        ws.cell(rr, 1, f.get("id", ""))
        ws.cell(rr, 2, f.get("factor", ""))
        ws.cell(rr, 3, "Yes" if f.get("present") else "No")
        sc = ws.cell(rr, 4, f.get("severity", ""))
        sc.fill = _sev_fill(f.get("severity", ""))
        ws.cell(rr, 5, f.get("likelihood", ""))
        ws.cell(rr, 6, f.get("existing_mitigation", ""))
        resid = ws.cell(rr, 7, f.get("residual_risk", ""))
        resid.fill = _sev_fill(f.get("residual_risk", ""))
        ws.cell(rr, 8, f.get("required_mitigation", ""))
        ws.cell(rr, 9, f.get("citation", ""))
        for c in range(1, 10):
            cell = ws.cell(rr, c)
            cell.border = BORDER
            cell.alignment = CENTER if c in (1, 3, 4, 5, 7) else WRAP
        rr += 1
    ws.freeze_panes = "A2"

    # ---- Feature Gating Matrix -----------------------------------------
    ws = wb.create_sheet("Feature Gating Matrix")
    cols = ["Feature", "Minimum age", "Threshold age", "Age assurance required?",
            "k-ID enforcement mechanism", "Current status", "Gap"]
    for i, c in enumerate(cols, 1):
        _hdr(ws.cell(1, i, c))
    widths = [28, 12, 13, 16, 30, 16, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    rr = 2
    for f in data.get("feature_matrix", []):
        ws.cell(rr, 1, f.get("feature", ""))
        ws.cell(rr, 2, f.get("min_age", ""))
        ws.cell(rr, 3, f.get("threshold_age", ""))
        ws.cell(rr, 4, f.get("assurance_required", ""))
        ws.cell(rr, 5, f.get("kid_mechanism", ""))
        sc = ws.cell(rr, 6, f.get("status", ""))
        # colour the status: ungated/missing = high, partial = med, done = none
        st = str(f.get("status", "")).lower()
        if any(k in st for k in ("ungated", "missing", "none", "not")):
            sc.fill = _sev_fill("high")
        elif any(k in st for k in ("partial", "claimed", "self")):
            sc.fill = _sev_fill("medium")
        elif any(k in st for k in ("gated", "done", "enforced", "ok")):
            sc.fill = _sev_fill("none")
        ws.cell(rr, 7, f.get("gap", ""))
        for c in range(1, 8):
            cell = ws.cell(rr, c)
            cell.border = BORDER
            cell.alignment = CENTER if c in (2, 3, 4) else WRAP
        rr += 1
    ws.freeze_panes = "A2"

    # ---- Obligations ----------------------------------------------------
    ws = wb.create_sheet("Obligations")
    cols = ["PP Tunas obligation", "Applies?", "Current status", "Gap",
            "Action", "k-ID mechanism", "Owner", "Citation"]
    for i, c in enumerate(cols, 1):
        _hdr(ws.cell(1, i, c))
    widths = [26, 9, 20, 26, 28, 24, 10, 32]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    rr = 2
    for o in data.get("obligations", []):
        ws.cell(rr, 1, o.get("requirement", ""))
        ws.cell(rr, 2, o.get("applies", ""))
        ws.cell(rr, 3, o.get("current_status", ""))
        ws.cell(rr, 4, o.get("gap", ""))
        ws.cell(rr, 5, o.get("action", ""))
        ws.cell(rr, 6, o.get("kid_mechanism", ""))
        ws.cell(rr, 7, o.get("owner", ""))
        ws.cell(rr, 8, o.get("citation", ""))
        for c in range(1, 9):
            cell = ws.cell(rr, c)
            cell.border = BORDER
            cell.alignment = CENTER if c in (2, 7) else WRAP
        rr += 1
    ws.freeze_panes = "A2"

    # ---- Sources --------------------------------------------------------
    ws = wb.create_sheet("Sources")
    cols = ["Topic", "Regulation", "Article", "Source URL", "neimo. verification"]
    for i, c in enumerate(cols, 1):
        _hdr(ws.cell(1, i, c))
    widths = [26, 26, 16, 52, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    rr = 2
    for s in data.get("sources", []):
        ws.cell(rr, 1, s.get("topic", ""))
        ws.cell(rr, 2, s.get("regulation", ""))
        ws.cell(rr, 3, s.get("article", ""))
        ws.cell(rr, 4, s.get("url", ""))
        ws.cell(rr, 5, s.get("verification", ""))
        for c in range(1, 6):
            ws.cell(rr, c).border = BORDER
            ws.cell(rr, c).alignment = WRAP
        rr += 1
    ws.freeze_panes = "A2"

    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="Build a PP Tunas risk register .xlsx")
    ap.add_argument("--in", dest="infile", required=True, help="Scored assessment JSON")
    ap.add_argument("--out", dest="outfile", required=True, help="Output .xlsx path")
    args = ap.parse_args()
    with open(args.infile, encoding="utf-8") as fh:
        data = json.load(fh)
    path = build(data, args.outfile)
    print(f"Wrote {path}")


if __name__ == "__main__":
    sys.exit(main())
